import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.views.decorators.csrf import csrf_exempt
from wh_former.models import Basket_master_data
from django.http import JsonResponse
from django.utils import timezone
from datetime import datetime
from django.contrib.auth.decorators import login_required

from django.utils.encoding import force_str
from django.utils.functional import Promise
from django.utils.html import strip_tags
import decimal
from datetime import date, datetime as dt


def to_xlsx_value(v):
    if isinstance(v, (list, tuple)) and v:
        v = v[0]

    if isinstance(v, Promise):
        v = force_str(v)
    if isinstance(v, str):
        v = strip_tags(v)
    if v is None:
        return ""
    if isinstance(v, decimal.Decimal):
        return float(v)

    if isinstance(v, (int, float, bool, date, dt, str)):
        return v

    return force_str(v)


import re
@login_required
def export_excel_button_api(request):
    payload = request.session.get('basket_export')
    if not payload:
        return JsonResponse({"error": "Không tìm thấy dữ liệu export trong session. Hãy tìm kiếm lại rồi xuất."}, status=404)

    headers = [to_xlsx_value(h) for h in payload.get("headers", [])]
    rows = payload.get("rows", [])

    if headers:
        headers.pop(-2)

    data = []
    for row in rows:
        data_item = []
        for idx, item in enumerate(row[0]):
            if idx == len(row[0]) - 2:
                continue
            elif idx == len(row[0]) - 3:
                text = re.sub(r"<.*?>", "", item[0]).strip()
                data_item.append(text)
            else:
                data_item.append(item[0])
        data.append(data_item)


    wb = Workbook()
    ws = wb.active
    ws.title = "Kết quả"

    font_body = Font(name='Cambria', size=14)
    font_header = Font(name='Cambria', size=14, bold=True)
    align_center = Alignment(horizontal='center', vertical='center')


    default_widths = [8, 26.6, 20, 13, 13, 20, 20]
    for i in range(1, len(headers) + 1):
        width = default_widths[i-1] if i-1 < len(default_widths) else 20
        ws.column_dimensions[get_column_letter(i)].width = width

    for i in range(1, len(data) + 2):
        ws.row_dimensions[i].height = 24

    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = font_header
        cell.alignment = align_center


    for idx, row in enumerate(data, start=2):
        ws.append(row)
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=idx, column=col)
            cell.font = font_body


    from io import BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    resp = HttpResponse(
        output.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp['Content-Disposition'] = f'attachment; filename="DN_basket_{datetime.today().strftime("%Y%m%d")}.xlsx"'
    return resp




#
#
#
# @csrf_exempt
# def export_excel_button_api(request):
#     if request.method == "POST":
#         payload = json.loads(request.body)
#         headers = payload.get('headers', [])
#         data = payload.get('data', [])
#         basket_no = payload.get('basket_no')
#
#
#         wb = Workbook()
#         ws = wb.active
#         ws.title = "Kết quả"
#
#         font_body = Font(name='Cambria', size=14)
#         font_header = Font(name='Cambria', size=14, bold=True)
#         align_center = Alignment(horizontal='center', vertical='center')
#
#         col_widths = [8, 26.6, 20, 13, 13, 20, 20][:len(headers)]
#         for i, width in enumerate(col_widths, start=1):
#             ws.column_dimensions[get_column_letter(i)].width = width
#
#         for i in range(1, len(data) + 2):
#             ws.row_dimensions[i].height = 24
#
#         ws.append(headers)
#         for col in range(1, len(headers) + 1):
#             cell = ws.cell(row=1, column=col)
#             cell.font = font_header
#             cell.alignment = align_center
#
#         for idx, row in enumerate(data, start=2):
#             ws.append(row)
#             for col in range(1, len(headers) + 1):
#                 cell = ws.cell(row=idx, column=col)
#                 cell.font = font_body
#
#         from io import BytesIO
#         output = BytesIO()
#         wb.save(output)
#         output.seek(0)
#         response = HttpResponse(output.read(),
#                                 content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
#
#         filename = f"DN_basket_{datetime.today().strftime('%Y%m%d')}.xlsx"
#         response['Content-Disposition'] = f'attachment; filename="{filename}"'
#
#         return response
#








@csrf_exempt
def basket_disable_button_api(request):
    if request.method == "POST":
        data = json.loads(request.body)
        basket_no = data.get('basket_no')
        try:
            obj = Basket_master_data.objects.get(basket_no=basket_no)
            # days_difference = (timezone.now().date() - obj.received_date).days
            # if days_difference >= 0:
            obj.is_active = 3
            obj.save()
            return JsonResponse({'status': 'success'})
            # else:
            #     return JsonResponse({'status': 'error', 'msg': f'Basket cannot be disabled. Only {days_difference} days since stock-in.'})
        except Basket_master_data.DoesNotExist:
            return JsonResponse({'status': 'not_found'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'msg': str(e)})